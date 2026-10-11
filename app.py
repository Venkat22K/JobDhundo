import os

import pandas as pd
from flask import Flask, render_template, request

from docx import Document
from pypdf import PdfReader

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from sentence_transformers import SentenceTransformer


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024


# ============================================================
# LOAD SBERT MODEL
# ============================================================

print("Loading SBERT model...")

sbert_model = SentenceTransformer(
    "all-MiniLM-L6-v2"
)

print("SBERT model loaded.")


# ============================================================
# TEXT PREPROCESSING
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = str(text).lower()

    import re
    import string

    # Remove email addresses
    text = re.sub(
        r"\S+@\S+",
        " ",
        text
    )

    # Remove URLs
    text = re.sub(
        r"http\S+|www\S+",
        " ",
        text
    )

    # Remove punctuation
    text = text.translate(
        str.maketrans(
            "",
            "",
            string.punctuation
        )
    )

    # Remove extra whitespace
    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# EXTRACT TEXT FROM DOCX
# ============================================================

def extract_docx(file):

    document = Document(file)

    text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text.append(
                paragraph.text.strip()
            )

    # Tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text.append(
                        cell.text.strip()
                    )

    return "\n".join(text)


# ============================================================
# EXTRACT TEXT FROM PDF
# ============================================================

def extract_pdf(file):

    reader = PdfReader(file)

    text = []

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:

            text.append(page_text)

    return "\n".join(text)


# ============================================================
# EXTRACT RESUME
# ============================================================

def extract_resume(file):

    filename = file.filename.lower()

    if filename.endswith(".pdf"):

        return extract_pdf(file)

    elif filename.endswith(".docx"):

        return extract_docx(file)

    elif filename.endswith(".txt"):

        return file.read().decode(
            "utf-8",
            errors="ignore"
        )

    else:

        raise ValueError(
            "Unsupported file type. "
            "Please upload PDF, DOCX, or TXT."
        )


# ============================================================
# TF-IDF MATCHING
# ============================================================

def calculate_tfidf(resume, job):

    documents = [
        resume,
        job
    ]

    vectorizer = TfidfVectorizer(
        lowercase=True
    )

    tfidf_matrix = vectorizer.fit_transform(
        documents
    )

    score = cosine_similarity(
        tfidf_matrix[0],
        tfidf_matrix[1]
    )[0][0]

    return float(score)


# ============================================================
# SBERT MATCHING
# ============================================================

def calculate_sbert(resume, job):

    resume_embedding = sbert_model.encode(
        [resume]
    )

    job_embedding = sbert_model.encode(
        [job]
    )

    score = cosine_similarity(
        resume_embedding,
        job_embedding
    )[0][0]

    return float(score)


# ============================================================
# HOME PAGE
# ============================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# ============================================================
# ANALYZE RESUME
# ============================================================

@app.route(
    "/analyze",
    methods=["POST"]
)
def analyze():

    resume_file = request.files.get(
        "resume"
    )

    job_description = request.form.get(
        "job_description",
        ""
    )

    # --------------------------------------------------------
    # Validate input
    # --------------------------------------------------------

    if not resume_file:

        return render_template(
            "index.html",
            error="Please upload a resume."
        )

    if not job_description.strip():

        return render_template(
            "index.html",
            error="Please enter a job description."
        )

    try:

        # ----------------------------------------------------
        # Extract resume
        # ----------------------------------------------------

        resume_text = extract_resume(
            resume_file
        )

        if not resume_text.strip():

            return render_template(
                "index.html",
                error="Could not extract text from the resume."
            )

        # ----------------------------------------------------
        # Clean text
        # ----------------------------------------------------

        cleaned_resume = clean_text(
            resume_text
        )

        cleaned_job = clean_text(
            job_description
        )

        # ----------------------------------------------------
        # Calculate scores
        # ----------------------------------------------------

        tfidf_score = calculate_tfidf(
            cleaned_resume,
            cleaned_job
        )

        sbert_score = calculate_sbert(
            cleaned_resume,
            cleaned_job
        )

        # ----------------------------------------------------
        # Display results
        # ----------------------------------------------------

        return render_template(
            "index.html",

            tfidf_score=round(
                tfidf_score,
                4
            ),

            sbert_score=round(
                sbert_score,
                4
            ),

            resume_name=resume_file.filename
        )

    except Exception as e:

        print("ERROR:", e)

        return render_template(
            "index.html",
            error=f"Error processing resume: {str(e)}"
        )


# ============================================================
# RUN APP
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=True
    )