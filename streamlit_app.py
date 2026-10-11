import firebase_admin
from firebase_admin import credentials, firestore
from datetime import datetime, timezone
import streamlit as st
import re
import string
import numpy as np

from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer


import firebase_admin
from firebase_admin import credentials, firestore

FIREBASE_PROJECT_ID = "jobdhundo-24f3b"



@st.cache_resource
def get_firestore_client():
    if not firebase_admin._apps:
        if "firebase" in st.secrets:
            # Use Streamlit Cloud Secrets
            firebase_config = dict(st.secrets["firebase"])
            cred = credentials.Certificate(firebase_config)
        else:
            # Use local Firebase key
            cred = credentials.Certificate("firebase-key.json")

        firebase_admin.initialize_app(
            cred,
            {"projectId": FIREBASE_PROJECT_ID}
        )

    return firestore.client()


st.set_page_config(
    page_title="JobDhundo",
    page_icon="💼",
    layout="wide"
)

st.title("💼 JobDhundo")
st.subheader("Find the right job. Match the right skills.")
st.write(
    "Compare your resume with a job description using "
    "TF-IDF keyword matching and SBERT semantic matching."
)


@st.cache_resource
def load_model():
    return SentenceTransformer("all-MiniLM-L6-v2")


def extract_text(uploaded_file):
    name = uploaded_file.name.lower()

    if name.endswith(".txt"):
        return uploaded_file.getvalue().decode("utf-8", errors="ignore")

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    if name.endswith(".docx"):
        document = Document(uploaded_file)
        paragraphs = [p.text for p in document.paragraphs]
        tables = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + tables)

    return ""


def clean_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


left, right = st.columns(2)

with left:
    st.markdown("### 📄 Upload Resume")
    resume_file = st.file_uploader(
        "Upload PDF, DOCX or TXT",
        type=["pdf", "docx", "txt"]
    )

with right:
    st.markdown("### 💼 Job Description")
    job_description = st.text_area(
        "Paste the complete job description",
        height=220,
        placeholder="Enter the role, required skills and responsibilities..."
    )

if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):

    if resume_file is None:
        st.error("Please upload a resume.")
    elif not job_description.strip():
        st.error("Please enter a job description.")
    else:
        try:
            resume_text = extract_text(resume_file)

            if not resume_text.strip():
                st.error("No readable text found in the resume.")
                st.stop()

            resume_clean = clean_text(resume_text)
            job_clean = clean_text(job_description)

            with st.spinner("Analyzing your resume..."):
                vectorizer = TfidfVectorizer(
                    stop_words="english"
                )

                vectors = vectorizer.fit_transform(
                    [resume_clean, job_clean]
                )

                tfidf_score = float(
                    cosine_similarity(vectors[0:1], vectors[1:2])[0][0]
                )

                model = load_model()

                embeddings = model.encode(
                    [resume_text, job_description],
                    normalize_embeddings=True
                )

                sbert_score = float(
                    np.dot(embeddings[0], embeddings[1])
                )
                try:
                    db = get_firestore_client()

                    db.collection("analyses").add({
                        "resume_filename": resume_file.name,
                        "job_description": job_description,
                        "tfidf_score": tfidf_score,
                        "sbert_score": sbert_score,
                        "created_at": datetime.now(timezone.utc)
                    })

                    st.success("Analysis saved to your history!")

                except Exception as e:
                    st.warning(
                        "Analysis completed, but saving to Firestore failed. "
                        f"Error: {e}"
                    )
            st.markdown("---")
            st.markdown("## 📊 Matching Results")

            col1, col2 = st.columns(2)

            with col1:
                st.markdown("### TF-IDF")
                st.metric("Keyword Similarity", f"{tfidf_score:.4f}")
                st.progress(max(0.0, min(1.0, tfidf_score)))

            with col2:
                st.markdown("### SBERT")
                st.metric("Semantic Similarity", f"{sbert_score:.4f}")
                st.progress(max(0.0, min(1.0, sbert_score)))

            difference = abs(tfidf_score - sbert_score)

            st.info(
                f"Absolute score difference: {difference:.4f}"
            )

            st.caption(
                "These are text similarity scores, not hiring probabilities "
                "or percentages of qualification. Scores from the two "
                "methods are not directly equivalent."
            )

        except Exception as error:
            st.error(f"Analysis failed: {error}")


st.markdown("---")
st.caption("JobDhundo | Resume–Job Matching using NLP")

st.markdown("---")
st.header("📚 Previous Analyses")

if st.button("Load History"):
    try:
        db = get_firestore_client()

        records = (
            db.collection("analyses")
            .order_by(
                "created_at",
                direction=firestore.Query.DESCENDING
            )
            .limit(20)
            .stream()
        )

        history = []

        for record in records:
            item = record.to_dict()
            history.append({
                "Resume": item.get("resume_filename", ""),
                "TF-IDF": round(item.get("tfidf_score", 0), 4),
                "SBERT": round(item.get("sbert_score", 0), 4),
                "Date": str(item.get("created_at", "")),
                "Job Description": item.get("job_description", "")
            })

        if history:
            st.dataframe(history, use_container_width=True)
        else:
            st.info("No saved analyses yet.")

    except Exception as e:
        st.error(f"Could not load history: {e}")