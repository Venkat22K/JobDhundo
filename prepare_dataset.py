import os
import re
import ast
import pandas as pd
from docx import Document


# ============================================================
# FILE PATHS
# ============================================================

CV_FOLDER = "data/raw/CV"
VACANCY_FILE = "data/raw/5_vacancies.csv"
ANNOTATION_FILE = "data/raw/annotations-for-the-first-30-vacancies.txt"
OUTPUT_FILE = "data/dataset.csv"


# ============================================================
# EXTRACT TEXT FROM DOCX
# ============================================================

def extract_docx(path):
    doc = Document(path)

    text = []

    # Normal paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text.strip())

    return "\n".join(text)


# ============================================================
# READ VACANCIES
# ============================================================

print("Reading vacancies...")

vacancies = pd.read_csv(VACANCY_FILE)

print("Vacancies loaded:", len(vacancies))


# ============================================================
# READ REFERENCE RANKINGS
# ============================================================

print("Reading reference rankings...")

with open(ANNOTATION_FILE, "r", encoding="utf-8") as f:
    annotation_text = f.read()


# Find the Annotator 2 section in the original file.
# We use this complete ranking set as the reference ranking.

match = re.search(
    r"ANNOTATOR_2_RANKINGS\s*=\s*(.*)",
    annotation_text,
    re.DOTALL
)

if not match:
    raise ValueError("Could not find reference rankings")


reference_text = match.group(1)


# Extract exactly five integers from each ranking array
reference_strings = re.findall(
    r"\[\s*\d+\s*,\s*\d+\s*,\s*\d+\s*,\s*\d+\s*,\s*\d+\s*\]",
    reference_text
)


reference_rankings = [
    ast.literal_eval(x)
    for x in reference_strings
]


# ============================================================
# CHECK REFERENCE RANKINGS
# ============================================================

print(
    "Reference rankings found:",
    len(reference_rankings)
)

if len(reference_rankings) != 30:
    raise ValueError(
        f"Expected 30 reference rankings, "
        f"but found {len(reference_rankings)}"
    )


# ============================================================
# CREATE DATASET
# ============================================================

print()
print("Creating dataset...")

rows = []


for cv_id in range(1, 31):

    # --------------------------------------------------------
    # Resume file
    # --------------------------------------------------------

    cv_path = os.path.join(
        CV_FOLDER,
        f"{cv_id}.docx"
    )

    if not os.path.exists(cv_path):

        print(f"WARNING: CV {cv_id} not found")

        resume_text = ""

    else:

        resume_text = extract_docx(cv_path)


    # --------------------------------------------------------
    # Check resume text
    # --------------------------------------------------------

    if not resume_text.strip():

        print(
            f"WARNING: CV {cv_id} has no extracted text"
        )


    # --------------------------------------------------------
    # Five vacancies
    # --------------------------------------------------------

    for vacancy_id in range(1, 6):

        vacancy_row = vacancies.iloc[
            vacancy_id - 1
        ]

        # Combine all vacancy columns
        job_description_parts = []

        for value in vacancy_row:

            if pd.notna(value):
                job_description_parts.append(
                    str(value)
                )

        job_description = "\n".join(
            job_description_parts
        )


        # ----------------------------------------------------
        # Reference ranking
        # ----------------------------------------------------

        reference_rank = (
            reference_rankings[cv_id - 1][
                vacancy_id - 1
            ]
        )


        # ----------------------------------------------------
        # Add row
        # ----------------------------------------------------

        rows.append({

            "cv_id": cv_id,

            "vacancy_id": vacancy_id,

            "resume_text": resume_text,

            "job_description": job_description,

            "reference_rank": reference_rank

        })


# ============================================================
# CREATE DATAFRAME
# ============================================================

df = pd.DataFrame(rows)


# ============================================================
# SAVE DATASET
# ============================================================

df.to_csv(
    OUTPUT_FILE,
    index=False,
    encoding="utf-8"
)


# ============================================================
# SUMMARY
# ============================================================

print()
print("========================================")
print("DATASET CREATED SUCCESSFULLY")
print("========================================")

print("Rows:", len(df))

print("Columns:", list(df.columns))

print("Output:", OUTPUT_FILE)


# ============================================================
# CHECK MISSING RESUMES
# ============================================================

missing = df[
    df["resume_text"].isna()
    |
    (df["resume_text"].str.strip() == "")
]

print()
print(
    "Missing resume rows:",
    len(missing)
)

if len(missing) > 0:

    print(
        "Problematic CVs:",
        missing["cv_id"].unique()
    )

else:

    print(
        "All resumes extracted successfully!"
    )


# ============================================================
# BASIC CHECKS
# ============================================================

print()

print(
    "Unique CVs:",
    df["cv_id"].nunique()
)

print(
    "Unique vacancies:",
    df["vacancy_id"].nunique()
)

print(
    "Dataset shape:",
    df.shape
)


# ============================================================
# ROW COUNT CHECK
# ============================================================

expected_rows = 30 * 5

if len(df) == expected_rows:

    print(
        f"Row count check: PASS ({len(df)} rows)"
    )

else:

    print(
        f"Row count check: FAILED "
        f"(expected {expected_rows}, got {len(df)})"
    )


print()
print("Dataset preparation complete.")